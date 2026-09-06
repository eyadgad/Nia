"""Refresh the NIA name and date normalization workbooks from PostgreSQL.

Run from the workspace root after starting the Compose postgres service:
    python export_normalization_workbooks.py
"""

from __future__ import annotations

import asyncio
import copy
import json
import os
from pathlib import Path
from typing import Any

import asyncpg
from openpyxl import load_workbook


DATABASE_URL = os.environ.get("ASSISTANT_MIGRATOR_DATABASE_URL", "")
FILES = {
    49: "Master List",
    91: "Confirmed",
    93: "Additional Deaths",
    94: "Potential",
}
WORKBOOKS = ("nia_name_normalization_all_lists.xlsx", "nia_date_normalization_all_lists.xlsx")


def source_value(payload: dict[str, Any], raw_keys: list[str]) -> Any:
    fields = payload.get("fields") or {}
    for key in raw_keys:
        value = fields.get(key)
        if isinstance(value, dict):
            value = value.get("raw") or value.get("display") or value.get("value")
        if value not in (None, ""):
            return value
    return None


def student_name(payload: dict[str, Any]) -> str | None:
    parts = (payload.get("canonical") or {}).get("name_parts") or {}
    name = " ".join(str(parts.get(key, "")).strip() for key in ("first", "middle", "last"))
    return name or None


def header_row(sheet) -> int:
    for row in range(1, min(sheet.max_row, 5) + 1):
        if sheet.cell(row, 1).value == "source_row_id":
            return row
    raise ValueError(f"{sheet.title} has no source_row_id header")


def clear_data(sheet, first_data_row: int) -> None:
    if first_data_row <= sheet.max_row:
        sheet.delete_rows(first_data_row, sheet.max_row - first_data_row + 1)


def write_rows(sheet, rows: list[dict[str, Any]], registry: dict[str, list[str]], *, date_book: bool) -> None:
    headers_at = header_row(sheet)
    headers = [sheet.cell(headers_at, column).value for column in range(1, sheet.max_column + 1)]
    first_data_row = headers_at + 2 if not date_book else headers_at + 2
    styles = [copy.copy(sheet.cell(first_data_row, column)._style) for column in range(1, sheet.max_column + 1)]
    clear_data(sheet, first_data_row)

    for row_number, record in enumerate(rows, first_data_row):
        payload = record["row_data_normalized"]
        canonical = payload.get("canonical") or {}
        name_parts = canonical.get("name_parts") or {}
        date_parts = canonical.get("date_parts") or {}
        for column, header in enumerate(headers, 1):
            value = None
            if header == "source_row_id":
                value = record["source_row_id"]
            elif date_book and header == "student":
                value = student_name(payload)
            elif date_book and isinstance(header, str) and header.endswith(" (source)"):
                field = header.removesuffix(" (source)")
                value = source_value(payload, registry.get(field, []))
            elif date_book and isinstance(header, str) and header.startswith("_"):
                field, _, component = header[1:].rpartition("_")
                value = (date_parts.get(field) or {}).get(component)
            elif not date_book and isinstance(header, str) and header.startswith("_"):
                value = name_parts.get(header[1:])
            elif not date_book and isinstance(header, str):
                value = source_value(payload, [header])
            cell = sheet.cell(row_number, column, value)
            cell._style = copy.copy(styles[column - 1])


async def fetch_data() -> tuple[dict[int, list[dict[str, Any]]], dict[int, dict[str, list[str]]]]:
    if not DATABASE_URL:
        raise RuntimeError("ASSISTANT_MIGRATOR_DATABASE_URL is not configured")
    connection = await asyncpg.connect(DATABASE_URL)
    try:
        records: dict[int, list[dict[str, Any]]] = {}
        registries: dict[int, dict[str, list[str]]] = {}
        for file_id in FILES:
            registry_rows = await connection.fetch(
                "SELECT semantic_field, raw_json_keys FROM assistant.field_registry WHERE file_id = $1",
                file_id,
            )
            registries[file_id] = {
                row["semantic_field"]: list(row["raw_json_keys"] or []) for row in registry_rows
            }
            rows = await connection.fetch(
                """
                SELECT n.source_row_id, n.row_data_normalized
                FROM public.file_data_normalized n
                JOIN public.file f ON f.id = n.file_id AND f.version = n.version
                WHERE n.file_id = $1 AND n.status = 'ready' AND COALESCE(f.is_delete, false) = false
                ORDER BY n.source_row_id
                """,
                file_id,
            )
            records[file_id] = []
            for row in rows:
                record = dict(row)
                payload = record["row_data_normalized"]
                record["row_data_normalized"] = json.loads(payload) if isinstance(payload, str) else payload
                records[file_id].append(record)
        return records, registries
    finally:
        await connection.close()


async def main() -> None:
    records, registries = await fetch_data()
    for workbook_name in WORKBOOKS:
        workbook = load_workbook(Path(workbook_name))
        date_book = workbook_name.startswith("nia_date")
        for file_id, sheet_name in FILES.items():
            write_rows(workbook[sheet_name], records[file_id], registries[file_id], date_book=date_book)
        workbook.save(workbook_name)
        print(f"wrote {workbook_name}")


if __name__ == "__main__":
    asyncio.run(main())
